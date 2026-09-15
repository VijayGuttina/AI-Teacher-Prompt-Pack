from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("build/commercial-reference.docx")


def set_font(style, name="Aptos", size=11, bold=False, colour=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-GB")
    lang.set(qn("w:eastAsia"), "en-GB")
    lang.set(qn("w:bidi"), "en-GB")
    if colour:
        style.font.color.rgb = RGBColor(*colour)


def prevent_word_breaking(style):
    ppr = style._element.get_or_add_pPr()

    # Disable automatic hyphenation for this paragraph style.
    suppress = ppr.find(qn("w:suppressAutoHyphens"))
    if suppress is None:
        suppress = OxmlElement("w:suppressAutoHyphens")
        ppr.append(suppress)
    suppress.set(qn("w:val"), "true")

    # Do NOT add w:wordWrap=false here. That setting controls Word's
    # East-Asian line-breaking behaviour and can cause undesirable character
    # wrapping in Latin text. Normal Word wrapping should handle whole words.


def add_field(paragraph, instruction):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


def configure_document_settings(doc):
    settings = doc.settings._element

    auto_hyphenation = settings.find(qn("w:autoHyphenation"))
    if auto_hyphenation is None:
        auto_hyphenation = OxmlElement("w:autoHyphenation")
        settings.append(auto_hyphenation)
    auto_hyphenation.set(qn("w:val"), "false")

    no_caps_hyphenation = settings.find(qn("w:doNotHyphenateCaps"))
    if no_caps_hyphenation is None:
        no_caps_hyphenation = OxmlElement("w:doNotHyphenateCaps")
        settings.append(no_caps_hyphenation)
    no_caps_hyphenation.set(qn("w:val"), "true")

    # Do not allow Word to compress characters to force lines into a box.
    spacing_control = settings.find(qn("w:characterSpacingControl"))
    if spacing_control is None:
        spacing_control = OxmlElement("w:characterSpacingControl")
        settings.append(spacing_control)
    spacing_control.set(qn("w:val"), "dontCompress")


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document_settings(doc)

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(20)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)

    styles = doc.styles
    normal = styles["Normal"]
    set_font(normal, size=11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    prevent_word_breaking(normal)

    for name, size in [("Title", 28), ("Subtitle", 15), ("Heading 1", 18), ("Heading 2", 16), ("Heading 3", 14)]:
        style = styles[name]
        set_font(style, name="Aptos Display" if name in ("Title", "Heading 1") else "Aptos", size=size, bold=True)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        prevent_word_breaking(style)

    custom = {
        "Prompt Number": (13, True, "Aptos"),
        "Prompt Title": (15, True, "Aptos"),
        "Label": (10, True, "Aptos"),
        "Prompt Text": (10.5, False, "Aptos"),
        "Teacher Tip": (10.5, False, "Aptos"),
        "Source Code": (9.5, False, "Consolas"),
    }
    for name, (size, bold, font_name) in custom.items():
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = styles["Normal"]
        set_font(style, name=font_name, size=size, bold=bold)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.05
        prevent_word_breaking(style)

    header = section.header.paragraphs[0]
    header.text = "AI Prompt Toolkit for Primary Teachers"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.style, size=8)
    for run in header.runs:
        run.font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Commercial Edition V1.0  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    add_field(footer, "PAGE")

    props = doc.core_properties
    props.title = "AI Prompt Toolkit for Primary Teachers"
    props.subject = "Commercial Edition V1.0"
    props.author = "Vijay Guttina"
    props.keywords = "primary teachers, ChatGPT, AI prompts, lesson planning, assessment"

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("AI Prompt Toolkit\nfor Primary Teachers")
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Commercial Edition V1.0")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("Plan faster. Teach better. Assess smarter. Save hours every week.")
    r.bold = True
    r.font.size = Pt(14)

    doc.add_page_break()
    doc.add_heading("Sample Chapter", level=1)
    doc.add_paragraph("This page demonstrates the formatting used by the commercial publication.")
    doc.add_heading("Prompt 001: Example prompt", level=2)

    for label, text in [
        ("Best for", "Turning a curriculum objective into a practical lesson."),
        ("Use when", "You have the learning goal but need a realistic first draft."),
        ("Difficulty", "Easy"),
    ]:
        p = doc.add_paragraph()
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(text)

    p = doc.add_paragraph()
    p.add_run("Copy and paste").bold = True
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Mm(160)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ("I teach [YEAR GROUP] pupils. The learning goal is [GOAL]. "
                 "Create a practical lesson with clear success criteria, a short retrieval task, "
                 "guided practice, independent practice and a final check for understanding. "
                 "Keep preparation realistic and make reasonable assumptions where details are missing.")
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6F8")
    tcPr.append(shd)

    p = doc.add_paragraph(style="Teacher Tip")
    p.add_run("Teacher tip: ").bold = True
    p.add_run("Give the assistant the actual teaching job, not just the topic.")
    doc.add_heading("Related prompts", level=3)
    doc.add_paragraph("Prompt 002: Example follow-up\nPrompt 003: Example assessment check")

    doc.save(OUT)
    print(f"Created {OUT}")


if __name__ == "__main__":
    main()
